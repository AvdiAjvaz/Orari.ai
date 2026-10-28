from io import BytesIO
from collections import Counter
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader
from app.main import app, School, verify
from app.examples import photo_example

client=TestClient(app)

def request_data():
    photo=photo_example();state=photo['state']
    return dict(school={**state,'teachers':[{'name':t['name']} for t in state['teachers']]},entries=photo['entries'],classes=state['classes'],school_name=state['schoolName'],title=state['scheduleTitle'],mode='all',content='teacher')

def test_photo_transcription_and_counts():
    response=client.get('/api/examples/photo');assert response.status_code==200
    photo=response.json();data=request_data();school=School(**data['school'])
    assert photo['verified'] and verify(school,photo['entries'])==[]
    assert len(photo['entries'])==199 and school.periods==7
    assert Counter(school.lessons[e['lesson']].classroom for e in photo['entries'])=={'10/8':33,'11/9':33,'12/9':33,'10/7':33,'11/8':34,'12/8':33}
    assert school.lessons[photo['entries'][0]['lesson']].teacher=='Meritë'
    assert school.lessons[photo['entries'][5]['lesson']].teacher=='Nazim'
    assert 'vendmbajtës' in photo['note']

def test_native_excel_keeps_all_days_and_cells():
    response=client.post('/api/export/xlsx',json=request_data());assert response.status_code==200
    book=load_workbook(BytesIO(response.content))
    assert len(book.worksheets)==5
    monday=book.worksheets[0]
    assert monday['B5'].value=='10/8' and monday['G5'].value=='12/8'
    assert monday['B6'].value=='Meritë' and monday['G6'].value=='Nazim'
    assert '2026/2027' in monday['A2'].value
    assert book.worksheets[4]['B12'].value=='—'

def test_native_word_has_all_day_tables():
    response=client.post('/api/export/docx',json=request_data());assert response.status_code==200
    doc=Document(BytesIO(response.content))
    assert len(doc.tables)==5
    assert doc.tables[0].cell(1,1).text=='Meritë'
    assert doc.tables[0].cell(1,6).text=='Nazim'
    assert any('Medreseja Alauddin' in p.text for p in doc.paragraphs)

def test_pdf_is_readable_with_albanian_text():
    response=client.post('/api/export/pdf',json=request_data());assert response.status_code==200
    assert response.content.startswith(b'%PDF')
    reader=PdfReader(BytesIO(response.content));text='\n'.join(p.extract_text() for p in reader.pages)
    assert 'Meritë' in text and 'E MËRKURË' in text and 'E PREMTE' in text
    assert 'Medreseja Alauddin' in text and 'prof. pa caktuar' in text

def test_export_revalidates_schedule_and_constraints():
    data=request_data();data['entries'][0]['slot']=data['entries'][1]['slot'];data['entries'][0]['lesson']=data['entries'][1]['lesson']
    assert client.post('/api/export/xlsx',json=data).status_code==422
    data=request_data();data['school']['teachers'][0]['max_hours']=0
    assert client.post('/api/export/pdf',json=data).status_code==422
    data=request_data();data['entries'][0]['lesson']=999
    assert client.post('/api/export/docx',json=data).status_code==422

def test_single_view_exports_and_excel_formula_safety():
    data=request_data();data.update(mode='classroom',value='10/8',school_name='=SUM(1,2)')
    response=client.post('/api/export/xlsx',json=data);assert response.status_code==200
    book=load_workbook(BytesIO(response.content));assert len(book.worksheets)==1
    assert book.active['A1'].data_type=='s'
    assert book.active['B6'].value.startswith('Lëndë e papërcaktuar\nMeritë')
    data['value']='nonexistent'
    assert client.post('/api/export/docx',json=data).status_code==422
