"""Native Office/PDF documents built from verified schedule projections."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, KeepTogether

DAYS = ['Hënë', 'Martë', 'Mërkurë', 'Enjte', 'Premte', 'Shtunë', 'Diel']
FONT_DIR = Path(__file__).parent / 'fonts'
pdfmetrics.registerFont(TTFont('Orari', str(FONT_DIR / 'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('OrariBold', str(FONT_DIR / 'DejaVuSans-Bold.ttf')))


def teacher_label(lesson):
    return lesson.subject + ' (prof. pa caktuar)' if lesson.teacher.startswith('Pa caktuar · ') else lesson.teacher


def projections(request):
    school = request.school
    index = {(e.slot, school.lessons[e.lesson].classroom): school.lessons[e.lesson] for e in request.entries}
    classes = list(dict.fromkeys(request.classes + [l.classroom for l in school.lessons]))
    if request.mode == 'all':
        output = []
        for day in range(school.days):
            rows = [['Ora'] + classes]
            for period in range(school.periods):
                row = [str(period + 1)]
                for classroom in classes:
                    lesson = index.get((day * school.periods + period, classroom))
                    text = '—'
                    if lesson:
                        text = lesson.subject if request.content == 'subject' else teacher_label(lesson)
                        if request.content == 'both':
                            text += '\n' + lesson.subject
                    row.append(text)
                rows.append(row)
            output.append(('E ' + DAYS[day].upper(), rows))
        return output
    field = request.mode
    if not any(getattr(l, field) == request.value for l in school.lessons):
        raise ValueError('Pamja e zgjedhur nuk përmban angazhime.')
    entries = {e.slot: school.lessons[e.lesson] for e in request.entries if getattr(school.lessons[e.lesson], field) == request.value}
    rows = [['Ora'] + DAYS[:school.days]]
    for period in range(school.periods):
        row = [str(period + 1)]
        for day in range(school.days):
            lesson = entries.get(day * school.periods + period)
            text = '—'
            if lesson:
                text = lesson.subject + '\n' + (teacher_label(lesson) if field == 'classroom' else lesson.classroom)
                if field == 'room':
                    text += '\n' + teacher_label(lesson)
                elif lesson.room:
                    text += '\n' + lesson.room
            row.append(text)
        rows.append(row)
    return [(request.value, rows)]


def document_bytes(request, format):
    tables = projections(request)
    school_name = request.school_name or 'Orari.ai'
    title = request.title or 'Orari javor mësimor'
    notes = []
    if any(l.subject == 'Lëndë e papërcaktuar' for l in request.school.lessons):
        notes.append('Lëndët që mungojnë në foto janë shënuar si të papërcaktuara.')
    if any(l.teacher.startswith('Pa caktuar · ') for l in request.school.lessons):
        notes.append('Sociologji: profesori duhet të caktohet; përkohësisht përdoret vendmbajtës për çdo klasë.')
    output = BytesIO()
    if format == 'xlsx':
        book = Workbook()
        book.remove(book.active)
        for heading, rows in tables:
            sheet = book.create_sheet(''.join(c for c in heading if c not in '[]:*?/\\')[:31] or 'Orari')
            columns = len(rows[0])
            for row_number, text in [(1, school_name), (2, title), (4, heading)]:
                sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=columns)
                sheet.cell(row_number, 1, text).data_type = 's'
                sheet.cell(row_number, 1).font = Font(size=14 if row_number != 2 else 11, bold=True, color='244C72')
            for r, row in enumerate(rows, 5):
                for c, text in enumerate(row, 1):
                    cell = sheet.cell(r, c, text)
                    cell.data_type = 's'  # Never interpret user input as an Excel formula.
                    cell.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center' if c == 1 or r == 5 else 'left')
                    cell.font = Font(size=10, bold=r == 5, color='172C31')
                    cell.fill = PatternFill('solid', fgColor='E7EFF8' if r == 5 or c == 1 else ('F6F8FB' if r % 2 else 'FFFFFF'))
                sheet.row_dimensions[r].height = 31 if any('\n' in text for text in row) else 24
            for c in range(1, columns + 1):
                sheet.column_dimensions[get_column_letter(c)].width = 6 if c == 1 else 24
            for note in notes:
                sheet.append([note])
                sheet.cell(sheet.max_row, 1).data_type = 's'
            sheet.freeze_panes = 'B6'
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = '1:5'
        book.save(output)
    elif format == 'docx':
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Inches(11.7), Inches(8.3)
        section.left_margin = section.right_margin = Inches(.45)
        doc.styles['Normal'].font.name = 'Calibri'
        doc.styles['Normal'].font.size = Pt(9)
        doc.add_heading(school_name, 0)
        doc.add_paragraph(title)
        for note in notes:
            doc.add_paragraph(note, style='Caption')
        for heading, rows in tables:
            doc.add_heading(heading, 2)
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Table Grid'
            for i, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
                    if i == 0:
                        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'E7EFF8'); cell._tc.get_or_add_tcPr().append(shading)
                props = table.rows[-1]._tr.get_or_add_trPr()
                props.append(OxmlElement('w:cantSplit'))
                if i == 0:
                    props.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph('')
        doc.save(output)
    else:
        # Wide schools use landscape, with automatic horizontal grouping to keep text legible.
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
        regular = ParagraphStyle('cell', fontName='Orari', fontSize=8, leading=10)
        heading_style = ParagraphStyle('heading', fontName='OrariBold', fontSize=13, leading=18, alignment=TA_CENTER, spaceAfter=8)
        story = [Paragraph(escape(school_name), heading_style), Paragraph(escape(title), heading_style)]
        for note in notes:
            story.append(Paragraph(escape(note), regular))
        story.append(Spacer(1, 12))
        for heading, rows in tables:
            for first in range(1, len(rows[0]), 8):
                subset = [[row[0]] + row[first:first+8] for row in rows]
                grid = [[Paragraph(escape(text).replace('\n', '<br/>'), regular) for text in row] for row in subset]
                cols = len(subset[0])
                table = Table(grid, colWidths=[28] + [(doc.width - 28) / (cols - 1)] * (cols - 1), repeatRows=1)
                table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .4, colors.HexColor('#C6D2DC')), ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E7EFF8')), ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F6F8FB')]), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5)]))
                label = heading + (f' · kolonat {first}–{min(first+7,len(rows[0])-1)}' if len(rows[0]) > 9 else '')
                story.extend([KeepTogether([Paragraph(escape(label), heading_style), table]), Spacer(1,16)])
        doc.build(story)
    return output.getvalue()
